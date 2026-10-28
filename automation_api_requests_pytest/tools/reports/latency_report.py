# tools/reports/latency_report.py
# 【工具】耗时报告 —— 拿一份提问集跑一遍，量出"每条要等多久"。
#
# 用法（在项目根目录下跑）：
#     .venv/Scripts/python.exe tools/reports/latency_report.py                 # 跑默认提问集
#     .venv/Scripts/python.exe tools/reports/latency_report.py --limit 5       # 只跑前 5 条，先看看
#     .venv/Scripts/python.exe tools/reports/latency_report.py --data xxx.yaml
#
# 会做三件事：跑一遍 → 控制台打表 → 导出 reports/assistant_latency_<时间>.xlsx
#
# ══ 为什么是"脚本"，不是"用例"？═════════════════════════════════════
# 因为它产出的是 **报告**，不是 **断言**。这两件事经常被混在一起：
#
#   · 用例的产出是红/绿 —— 回答"有没有退化"，用来守门（CI 里拦人）；
#   · 报告的产出是**数字** —— 回答"现在有多慢、慢在哪"，用来做决策。
#
# 把"耗时不超过 3 秒"写成 assert 是很多人第一反应会做的事，但它会**随机变红**：
# 同一段代码，机器在跑编译或杀毒软件在扫描时就超标了。而随机变红的套件比没有套件更糟
# ——没人再信红色，真出事那天也被当成噪音忽略掉。
# 所以这里只测、只打印、只导出，**一个断言都不做**。
#
# ══ 指标口径（对照 AI学伴评测方案的「性能指标」那节）══════════════════
# 那边是 SSE 流式接口，能测四个数：意图识别耗时 / RAG检索耗时 / 首字符耗时 / 总耗时。
# 我们这边**只有总耗时**，原因很实在：
#   /assistant/chat 是同步返回的 —— 回答全部生成完才一次性吐回来，
#   "第一个字符到"和"最后一个字符到"是同一时刻，中间那段测不出来。
#   想补"首字符耗时"，得先让后端把这个接口换成流式。
#   这是**能力的边界**，不是偷懒 —— 报告里标清楚，比编一个假指标强。

import argparse
import sys
from datetime import datetime
from pathlib import Path

# 脚本放在 tools/reports/ 下，而 api/ common/ config/ 都在**上两层**。
# 把项目根塞进搜索路径，这样不管从哪个目录调用都能 import 到（双击、IDE 里都行）。
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from api import Api                              # noqa: E402
from common.data_loader import load_cases        # noqa: E402
from common.logger import log, pad               # noqa: E402
from config.config import TEST_USER              # noqa: E402

REPORT_DIR = ROOT / "reports"
DEFAULT_DATA = "assistant_intent_cases.yaml"
SNIPPET_CHARS = 48        # 表格里每条回答只留一小段，够看出"答的是什么方向"就行


def percentile(values, p):
    """取第 p 百分位（最近秩法，够用）。

    P95 比"平均值"更该看：平均值会被大量快请求稀释，
    而用户体验是被**最慢的那几条**决定的 —— 一次等 8 秒就够劝退了。
    """
    ordered = sorted(values)
    if not ordered:
        return 0.0
    rank = round(p / 100 * (len(ordered) - 1))
    return ordered[max(0, min(len(ordered) - 1, rank))]


def measure(api, headers, cases):
    """跑一遍提问集，逐条量耗时。返回 list[dict]。

    单条失败**不中断**，记一行错误继续跑 —— 报告要的是"全貌"，
    跑到第 3 条就崩掉，剩下 12 条的情况你永远看不到。
    （这一点和用例相反：用例失败必须立刻停下来告诉你哪条错了。）
    """
    results = []
    for index, case in enumerate(cases, 1):
        resp = api.assistant.chat({"rawPrompt": case["question"]}, headers=headers)
        ms = resp.elapsed.total_seconds() * 1000

        if resp.status_code != 200:
            # 失败也记耗时 —— "慢到超时"本身就是要看的现象
            results.append({
                "name": case["name"], "intent": case.get("intent_cn", ""),
                "question": case["question"], "ms": ms,
                "answer": f"[失败] HTTP {resp.status_code} {resp.text[:80]}",
            })
            continue

        data = resp.json()["data"]
        results.append({
            "name": case["name"], "intent": case.get("intent_cn", ""),
            "question": case["question"], "ms": ms,
            "answer": " ".join(data["assistantMessage"]["text"].split()),
        })

        # 每聊一句都会在库里建会话，不留垃圾：量完就删
        api.assistant.delete_conversation(data["conversationId"], headers=headers)

        print(f"  {index:>3}/{len(cases)}  {pad(case['name'], 22)} {ms:>7.0f}ms")

    return results


def print_table(results):
    """控制台打一张表 —— 顺序就是提问集顺序，方便和数据集对着看。"""
    print()
    log.info("· 耗时明细（单位 ms）")
    for r in results:
        log.info("  %s %s %s %6.0f   %s",
                 pad(r["name"], 22), pad(r["intent"], 12), pad(r["question"], 24),
                 r["ms"], r["answer"][:SNIPPET_CHARS])


def print_summary(results):
    """汇总。分两组打：整体一组，**按意图拆开**一组。

    按意图拆开是有用的 —— 它能回答"是整体都慢，还是某一类特别慢"。
    如果几十条整体平均 4s、但"越界类"单独 9s，那问题多半在"模型要先拒绝再解释"
    这个动作上，而不是模型本身慢。整体平均值会把这种差异抹平。
    """
    ms_list = [r["ms"] for r in results]
    log.info("· 汇总  条数=%d  平均=%.0fms  P50=%.0fms  P95=%.0fms  最快=%.0fms  最慢=%.0fms  总耗时=%.1fs",
             len(ms_list), sum(ms_list) / len(ms_list), percentile(ms_list, 50),
             percentile(ms_list, 95), min(ms_list), max(ms_list),
             sum(ms_list) / 1000)

    by_intent = {}
    for r in results:
        by_intent.setdefault(r["intent"], []).append(r["ms"])
    for intent, values in by_intent.items():
        log.info("  %s 条数=%d  平均=%.0fms  最慢=%.0fms",
                 pad(intent or "(未标意图)", 12), len(values),
                 sum(values) / len(values), max(values))


def write_excel(results, path):
    """导出两张表：明细 + 汇总。

    为什么导出 Excel 而不是只打日志？—— 因为报告是**给人看、要发出去**的，
    日志是给排查用的。两张表分工：
      · 明细 —— 每条一问一答加耗时，用来挑"最慢的几条"去看它答了什么；
      · 汇总 —— 一个数，用来写进周报和对比"改之前 / 改之后"。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["序号", "用例名", "意图", "提问", "耗时(ms)", "回答摘要"])
    for i, r in enumerate(results, 1):
        ws.append([i, r["name"], r["intent"], r["question"],
                   round(r["ms"]), r["answer"][:200]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEF", (6, 22, 12, 28, 10, 80)):
        ws.column_dimensions[col].width = width

    ms_list = [r["ms"] for r in results]
    ws2 = wb.create_sheet("汇总")
    ws2.append(["指标", "值"])
    for key, value in [
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("被测接口", "POST /assistant/chat"),
        ("条数", len(ms_list)),
        ("平均耗时(ms)", round(sum(ms_list) / len(ms_list))),
        ("P50(ms)", round(percentile(ms_list, 50))),
        ("P95(ms)", round(percentile(ms_list, 95))),
        ("最快(ms)", round(min(ms_list))),
        ("最慢(ms)", round(max(ms_list))),
        ("总耗时(s)", round(sum(ms_list) / 1000, 1)),
        ("未测指标", "意图识别耗时 / RAG耗时 / 首字符耗时 —— 接口是同步返回，测不出来"),
    ]:
        ws2.append([key, value])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 72

    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="AI 助手提问集耗时报告")
    parser.add_argument("--data", default=DEFAULT_DATA,
                        help=f"data/ 下的提问集文件名（默认 {DEFAULT_DATA}）")
    parser.add_argument("--limit", type=int, default=None,
                        help="只跑前 N 条 —— 数据集一大就得先试跑，别一上来跑全量")
    parser.add_argument("--out", default=None, help="导出路径（默认按时间自动命名）")
    args = parser.parse_args()

    cases = load_cases(args.data)
    if args.limit:
        cases = cases[:args.limit]

    log.info("· 耗时报告开始  提问集=%s  条数=%d", args.data, len(cases))

    api = Api()
    token = api.assert_ok(api.auth.login(TEST_USER))["data"]["token"]
    headers = {"Authorization": "Bearer " + token}

    results = measure(api, headers, cases)
    print_table(results)
    print_summary(results)

    path = Path(args.out) if args.out else (
        REPORT_DIR / f"assistant_latency_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    write_excel(results, path)
    log.info("· 报告已导出  %s", path)


if __name__ == "__main__":
    # Windows 控制台默认不是 UTF-8，中文一律会变成乱码 —— 不崩就行（文件里是好的）。
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    main()
