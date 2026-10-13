# common/data_loader.py
# 【数据驱动 DDT 的"读数据"零件】
#
# 目标：让测试代码**只关心"怎么测"，不关心"数据存在哪、什么格式"**。
# 用例那边永远只拿到 list[dict]（每行一个字典），换数据源不改用例。
#
# 类比：用例是"厨师"，数据文件是"食材箱"。厨师不该管食材是从冰箱(JSON)、
#       纸箱(CSV)、保鲜盒(YAML) 还是 Excel 表里拿的，他只管"拿到食材就下锅"。
#
# 已支持格式：.json / .csv / .yaml(.yml) / .xlsx

import csv
import io
import json
from pathlib import Path

import yaml

# 数据目录：项目根下的 data/（本文件在 common/ 里，所以往上一级）
DATA_DIR = Path(__file__).resolve().parent.parent / "data"


def _read_text(path):
    """按常见编码依次尝试读取整个文本文件。

    为什么不能直接 utf-8？因为 Excel 在中文 Windows 上存 CSV 有两个经典坑：
      · UTF-8 **带 BOM**：文件头会多出 \\ufeff，害得第一列表头变成 "\\ufeffname"
      · GBK / ANSI：压根不是 UTF-8
    `utf-8-sig` 会自动吃掉 BOM（没 BOM 也正常工作）；不行再退回 gbk。
    这是读"别人给的数据文件"时的标准防御姿势。
    """
    for enc in ("utf-8-sig", "gbk"):
        try:
            with path.open(encoding=enc, newline="") as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法解码数据文件（已尝试 utf-8-sig / gbk）：{path}")


# sheet="*" 时，每行会多出这个键，值是它来自哪张工作表。
# 定义成常量而不是到处写字符串字面量 —— 用例那边 `row[SHEET_KEY]` 比 `row["sheet"]` 好找。
SHEET_KEY = "_sheet"

# sheet 参数的"读全部工作表"取值
ALL_SHEETS = "*"


def _read_excel_sheet(ws):
    """读**一张**工作表：第一行当表头，其余行当数据。空单元格统一成空字符串。"""
    rows_iter = ws.iter_rows(values_only=True)
    header = ["" if h is None else str(h).strip() for h in next(rows_iter)]
    out = []
    for r in rows_iter:
        if r is None or all(v is None for v in r):
            continue                          # 跳过全空行（Excel 底部常见的幽灵行）
        row = {}
        for i, key in enumerate(header):
            v = r[i] if i < len(r) else None
            row[key] = "" if v is None else v
        out.append(row)
    return out


def _read_excel(path, sheet=None):
    """读 .xlsx。sheet 取三种值：

      · None（默认）—— 只读**活动工作表**。单表文件照旧，老调用一个字不用改。
      · "*"        —— 读**全部工作表**，每行多一个 `_sheet` 键，记着它来自哪张表。
      · "表名"      —— 只读指定的那张。

    **为什么要能一次读全部？**
    因为这能让"**工作表名**"本身承载信息 —— 也就是把**期望结果编码进表名**：

        工作表「重难点-课节级」  → 这一表的用例，都期望被识别成"重难点"意图
        工作表「出题-单选题」    → 这一表的用例，都期望被识别成"出题"意图

    这样做的好处是：**加一批样本 = 加一张工作表**，用例代码一行都不用改；
    而且"期望值"和"样本"分开放 —— 改判定规则只动表名，不碰任何一条数据。
    （这套做法照搬自 AI学伴的意图识别评测方案，见 docs/AI助手模块-迭代记录.md 迭代 4。）

    表的顺序按文件里的 sheet 顺序，不按名字排序 —— 保持和人在 Excel 里看到的一致。
    """
    from openpyxl import load_workbook  # 按需导入，不用 Excel 的人不必装

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None:
            return _read_excel_sheet(wb.active)
        if sheet == ALL_SHEETS:
            out = []
            for name in wb.sheetnames:
                for row in _read_excel_sheet(wb[name]):
                    out.append({**row, SHEET_KEY: name})
            return out
        if sheet not in wb.sheetnames:
            raise ValueError(
                f"工作表不存在：{sheet!r}；{path.name} 里只有 {wb.sheetnames}")
        return _read_excel_sheet(wb[sheet])
    finally:
        wb.close()


def load_cases(filename, int_fields=(), sheet=None):
    """从 data/ 目录读取用例数据，按扩展名自动选解析方式。

    参数：
        filename   —— data/ 下的文件名，如 "login_cases.yaml" / "login_cases.xlsx"
        int_fields —— 需要转成整数的列名（CSV 读出来全是字符串，见下）
        sheet      —— 只对 .xlsx 有意义：None=活动表（默认）/ "*"=全部表 / "表名"。
                      传 "*" 时每行会多一个 `_sheet` 键（见 SHEET_KEY）记着表名。
                      非 Excel 文件传了它 → 直接报错，别让人以为生效了却没生效。

    返回：list[dict]，每个 dict 就是一条用例的数据。
    """
    path = DATA_DIR / filename
    if not path.exists():
        raise FileNotFoundError(f"找不到数据文件：{path}")

    suffix = path.suffix.lower()

    if sheet is not None and suffix != ".xlsx":
        # 宁可报错也不要静默忽略：传了 sheet 却读到全部数据，是最难查的那种错
        raise ValueError(f"sheet 参数只对 .xlsx 有意义，{path.name} 是 {suffix} 文件")

    if suffix == ".json":
        rows = json.loads(_read_text(path))               # JSON 自带类型：数字就是数字
    elif suffix == ".csv":
        rows = list(csv.DictReader(io.StringIO(_read_text(path))))  # CSV 坑：读出来**全是字符串**
    elif suffix in (".yaml", ".yml"):
        rows = yaml.safe_load(_read_text(path))           # YAML 自带类型，且能写注释
    elif suffix == ".xlsx":
        rows = _read_excel(path, sheet=sheet)             # Excel：数字列一般是真数字
    else:
        raise ValueError(f"暂不支持的数据格式：{suffix}（老式 .xls 请先另存为 .xlsx）")

    if not isinstance(rows, list):
        raise ValueError(f"数据文件格式不对，顶层应为列表：{path}")

    # 统一把指定列转成 int，别让每个用例都去操心类型
    for row in rows:
        for key in int_fields:
            if row.get(key) not in (None, ""):
                row[key] = int(row[key])
    return rows
